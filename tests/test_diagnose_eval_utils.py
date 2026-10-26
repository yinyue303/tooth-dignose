import hashlib
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

try:
    from models.diagnose_eval_utils import (
        calculate_grouped_metrics,
        classify_language,
        format_metric_summary,
        normalize_metric_text,
        parse_image_paths,
        prompt_for_language,
        tokenize_chinese,
        validate_aligned_lengths,
        validate_language_values,
        validate_required_columns,
    )
except ImportError as exc:
    IMPORT_ERROR = exc
else:
    IMPORT_ERROR = None


class DiagnoseEvalUtilsTest(unittest.TestCase):
    def setUp(self):
        if IMPORT_ERROR is not None:
            self.fail(f"诊断评估辅助模块尚不可用：{IMPORT_ERROR}")

    def test_language_classification_only_needs_question_text(self):
        self.assertEqual(classify_language("患者右下后牙疼痛"), 0)
        self.assertEqual(classify_language("Patient reports tooth pain."), 1)
        self.assertEqual(classify_language("K07.303：修复体复诊"), 0)

    def test_language_validation_rejects_missing_and_unknown_values(self):
        self.assertEqual(validate_language_values([0, 1, 0]), [0, 1, 0])
        for values in ([0, None], [0, ""], [0, 2], [True, 1]):
            with self.subTest(values=values):
                with self.assertRaisesRegex(ValueError, "language"):
                    validate_language_values(values)

    def test_required_columns_are_checked(self):
        validate_required_columns(["id", "q", "a", "images", "language"])
        with self.assertRaisesRegex(ValueError, "language"):
            validate_required_columns(["id", "q", "a", "images"])

    def test_prompt_is_selected_per_sample_language(self):
        self.assertEqual(prompt_for_language(0, "中文", "English"), "中文")
        self.assertEqual(prompt_for_language(1, "中文", "English"), "English")
        with self.assertRaisesRegex(ValueError, "language"):
            prompt_for_language(3, "中文", "English")

    def test_metric_normalization_handles_compatibility_width_case_and_space(self):
        self.assertEqual(
            normalize_metric_text("  ＡＢＣ\u3000Ｋ０７．３０３  "),
            "abc k07.303",
        )

    def test_chinese_tokenizer_preserves_codes_numbers_and_unicode(self):
        self.assertEqual(
            tokenize_chinese("牙位４７ K07.303，㐀﨑〇🙂"),
            ["牙", "位", "47", "k07.303", ",", "㐀", "﨑", "〇", "🙂"],
        )

    def test_empty_text_has_no_tokens(self):
        self.assertEqual(tokenize_chinese(" \u3000\n\t"), [])

    def test_image_paths_are_parsed_without_executing_code(self):
        self.assertEqual(
            parse_image_paths("['images/a.jpg', '/images/b.png']"),
            ["images/a.jpg", "/images/b.png"],
        )
        for value in (
            "__import__('os').system('echo unsafe')",
            "('images/a.jpg',)",
            "['images/a.jpg', 3]",
        ):
            with self.subTest(value=value):
                with self.assertRaisesRegex(ValueError, "images"):
                    parse_image_paths(value)

    def test_grouped_metrics_keep_languages_separate_and_skip_chinese_meteor(self):
        references = ["牙47", "alpha beta", "病历", "gamma"]
        predictions = ["牙47", "alpha", "诊断", "gamma"]
        languages = [0, 1, 0, 1]

        def bleu_scorer(reference_tokens, prediction_tokens):
            return 1.0 if reference_tokens == prediction_tokens else 0.25

        def meteor_scorer(reference_tokens, prediction_tokens):
            if any("牙" in token or "病" in token for token in reference_tokens):
                raise AssertionError("中文样本不应调用 METEOR")
            return 0.75 if reference_tokens == prediction_tokens else 0.5

        metrics = calculate_grouped_metrics(
            references,
            predictions,
            languages,
            [0.2, 0.8, 0.4, 0.6],
            bleu_scorer,
            meteor_scorer,
            lambda text: normalize_metric_text(text).split(),
        )

        self.assertEqual(metrics["english_count"], 2)
        self.assertAlmostEqual(metrics["english_bleu"], 0.625)
        self.assertAlmostEqual(metrics["english_meteor"], 0.625)
        self.assertAlmostEqual(metrics["english_cosine"], 0.7)
        self.assertEqual(metrics["chinese_count"], 2)
        self.assertAlmostEqual(metrics["chinese_bleu"], 0.625)
        self.assertAlmostEqual(metrics["chinese_cosine"], 0.3)
        self.assertNotIn("chinese_meteor", metrics)

    def test_empty_reference_or_prediction_scores_zero_without_calling_scorers(self):
        def unexpected_scorer(_reference_tokens, _prediction_tokens):
            raise AssertionError("空文本不应调用指标函数")

        metrics = calculate_grouped_metrics(
            ["", "English"],
            ["中文", ""],
            [0, 1],
            [0.1, 0.2],
            unexpected_scorer,
            unexpected_scorer,
            lambda text: text.split(),
        )

        self.assertEqual(metrics["chinese_bleu"], 0.0)
        self.assertEqual(metrics["english_bleu"], 0.0)
        self.assertEqual(metrics["english_meteor"], 0.0)

    def test_alignment_and_nonempty_groups_are_required(self):
        with self.assertRaisesRegex(ValueError, "数量"):
            validate_aligned_lengths(["a"], ["b", "c"], [0], [0.1])
        with self.assertRaisesRegex(ValueError, "中文"):
            calculate_grouped_metrics(
                ["a"],
                ["a"],
                [1],
                [1.0],
                lambda _r, _p: 1.0,
                lambda _r, _p: 1.0,
                lambda text: text.split(),
            )

    def test_metric_summary_has_only_requested_grouped_metrics(self):
        output = format_metric_summary(
            {
                "english_count": 2,
                "english_bleu": 0.1,
                "english_meteor": 0.2,
                "english_cosine": 0.3,
                "chinese_count": 3,
                "chinese_bleu": 0.4,
                "chinese_cosine": 0.5,
            }
        )
        self.assertEqual(
            output,
            "English samples: 2\n"
            "English BLEU: 0.100000\n"
            "English METEOR: 0.200000\n"
            "English Cosine Similarity: 0.300000\n"
            "Chinese samples: 3\n"
            "Chinese BLEU: 0.400000\n"
            "Chinese Cosine Similarity: 0.500000",
        )
        self.assertNotIn("Chinese METEOR", output)


class DiagnosticWorkbookTest(unittest.TestCase):
    def test_language_column_is_complete_without_changing_original_cells(self):
        workbook_path = ROOT / "dataset" / "excel" / "test-diagnostic.xlsx"
        workbook = load_workbook(workbook_path, data_only=False)
        self.assertEqual(workbook.sheetnames, ["data"])
        sheet = workbook["data"]

        original_cells = hashlib.sha256()
        for row in sheet.iter_rows(min_col=1, max_col=4):
            for cell in row:
                original_cells.update(repr(cell.value).encode("utf-8"))
                original_cells.update(b"\0")

        self.assertEqual(
            original_cells.hexdigest(),
            "ba013e59ade3d194ffafe31cfbd8c3576b828c4b7fd1a92bba1aa785c00f4ed6",
        )
        self.assertEqual(sheet.max_row, 1800)
        self.assertEqual(
            [sheet.cell(1, column).value for column in range(1, 6)],
            ["id", "q", "a", "images", "language"],
        )
        languages = [sheet.cell(row, 5).value for row in range(2, 1801)]
        self.assertEqual(languages.count(0), 917)
        self.assertEqual(languages.count(1), 882)
        self.assertEqual(set(languages), {0, 1})
        self.assertEqual(sheet.tables["DiagnosticData"].ref, "A1:E1800")
        self.assertEqual(sheet.cell(1, 5).style_id, sheet.cell(1, 4).style_id)
        self.assertEqual(sheet.cell(1800, 5).style_id, sheet.cell(1800, 4).style_id)


if __name__ == "__main__":
    unittest.main()
