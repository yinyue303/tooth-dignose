import hashlib
import os
import shutil
import sys
import tempfile
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.table import TableColumn


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from models.diagnose_eval_utils import classify_language


WORKBOOK_PATH = ROOT / "dataset" / "excel" / "test-diagnostic.xlsx"
EXPECTED_HEADERS = ["id", "q", "a", "images"]
EXPECTED_ROWS = 1800
EXPECTED_COUNTS = {0: 917, 1: 882}


def cell_value_checksum(sheet, max_column=4):
    digest = hashlib.sha256()
    for row in sheet.iter_rows(min_col=1, max_col=max_column):
        for cell in row:
            digest.update(repr(cell.value).encode("utf-8"))
            digest.update(b"\0")
    return digest.hexdigest()


def update_workbook(source, destination):
    shutil.copy2(source, destination)
    workbook = load_workbook(destination, data_only=False)
    if workbook.sheetnames != ["data"]:
        raise ValueError(f"工作表结构不符合预期：{workbook.sheetnames!r}")

    sheet = workbook["data"]
    if sheet.max_row != EXPECTED_ROWS:
        raise ValueError(f"数据行数不符合预期：{sheet.max_row - 1}")
    headers = [sheet.cell(1, column).value for column in range(1, 5)]
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"原始列不符合预期：{headers!r}")

    original_checksum = cell_value_checksum(sheet)
    sheet.cell(1, 5, "language")
    sheet.cell(1, 5)._style = copy(sheet.cell(1, 4)._style)
    sheet.cell(1, 5).alignment = copy(sheet.cell(1, 4).alignment)

    counts = {0: 0, 1: 0}
    for row in range(2, sheet.max_row + 1):
        language = classify_language(sheet.cell(row, 2).value)
        language_cell = sheet.cell(row, 5, language)
        language_cell._style = copy(sheet.cell(row, 4)._style)
        language_cell.alignment = copy(sheet.cell(row, 4).alignment)
        counts[language] += 1

    sheet.column_dimensions["E"].width = 12
    table = sheet.tables["DiagnosticData"]
    table.ref = f"A1:E{sheet.max_row}"
    if table.autoFilter is not None:
        table.autoFilter.ref = table.ref
    if len(table.tableColumns) == 4:
        table.tableColumns.append(TableColumn(id=5, name="language"))

    if counts != EXPECTED_COUNTS:
        raise ValueError(f"language 计数不符合预期：{counts!r}")
    if cell_value_checksum(sheet) != original_checksum:
        raise ValueError("写入 language 时原始 A:D 单元格发生变化")

    workbook.save(destination)
    return original_checksum


def verify_workbook(path, original_checksum):
    workbook = load_workbook(path, data_only=False)
    if workbook.sheetnames != ["data"]:
        raise ValueError("保存后工作表名称或数量发生变化")
    sheet = workbook["data"]
    if cell_value_checksum(sheet) != original_checksum:
        raise ValueError("保存后原始 A:D 单元格发生变化")
    if [sheet.cell(1, column).value for column in range(1, 6)] != [
        "id",
        "q",
        "a",
        "images",
        "language",
    ]:
        raise ValueError("保存后列名不符合预期")
    languages = [sheet.cell(row, 5).value for row in range(2, sheet.max_row + 1)]
    counts = {0: languages.count(0), 1: languages.count(1)}
    if counts != EXPECTED_COUNTS or set(languages) != {0, 1}:
        raise ValueError(f"保存后 language 数据无效：{counts!r}")
    if sheet.tables["DiagnosticData"].ref != "A1:E1800":
        raise ValueError("保存后表格范围未扩展到 language 列")


def main():
    handle, temporary_name = tempfile.mkstemp(
        prefix="test-diagnostic.", suffix=".tmp.xlsx", dir=WORKBOOK_PATH.parent
    )
    os.close(handle)
    temporary_path = Path(temporary_name)
    try:
        original_checksum = update_workbook(WORKBOOK_PATH, temporary_path)
        verify_workbook(temporary_path, original_checksum)
        os.replace(temporary_path, WORKBOOK_PATH)
    finally:
        temporary_path.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
